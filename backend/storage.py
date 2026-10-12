"""Local storage and session state."""

from __future__ import annotations

import json
import logging
import re
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import UploadFile

from .config import Settings
from .rag_index import LocalEmbeddingIndex
from .foundation import SessionRepository, contained_path, validate_identifier
from .auth import principal_id
from fastapi import HTTPException

logger = logging.getLogger("data_extraction.storage")


class SessionStore:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.root = settings.local_storage_dir
        self.upload_root = self.root / "uploads"
        self.index_root = self.root / "indexes"
        self.embedding_index = LocalEmbeddingIndex(settings.local_embedding_dir.expanduser())
        self.session_store = settings.session_store
        self.root.mkdir(parents=True, exist_ok=True)
        self.upload_root.mkdir(parents=True, exist_ok=True)
        self.index_root.mkdir(parents=True, exist_ok=True)
        self.session_store.parent.mkdir(parents=True, exist_ok=True)
        self.repository = SessionRepository(self.session_store.with_suffix(".sqlite3"))
        self.repository.migrate_json(self.session_store)
        self.sessions, self.revisions = self.repository.load()
        self._saved = {key: json.dumps(value, sort_keys=True) for key, value in self.sessions.items()}

    def _load(self) -> Dict[str, Dict[str, Any]]:
        if not self.session_store.exists():
            return {}
        try:
            return json.loads(self.session_store.read_text(encoding="utf-8"))
        except Exception:
            return {}

    def save(self) -> None:
        for key, value in self.sessions.items():
            encoded = json.dumps(value, sort_keys=True)
            if self._saved.get(key) != encoded:
                self.revisions[key] = self.repository.write(key, value, self.revisions.get(key, 0))
                self._saved[key] = encoded

    def get_or_create(self, session_id: str) -> Dict[str, Any]:
        validate_identifier(session_id)
        owner = principal_id.get()
        if not owner:
            raise HTTPException(401, "Authentication required")
        if session_id not in self.sessions:
            self.sessions[session_id] = {
                "session_id": session_id,
                "user_id": owner,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "uploaded_files": [],
                "table_data": {"columns": [], "rows": []},
            }
            self.save()
            logger.info("session_store.created session_id=%s session_store=%s", session_id, self.session_store)
        session = self.sessions[session_id]
        if session.get("user_id") != owner:
            raise HTTPException(403, "Session access denied; legacy ownership requires migration")
        return session

    def get(self, session_id: str) -> Optional[Dict[str, Any]]:
        return self.sessions.get(session_id)

    async def save_uploads(self, session_id: str, files: List[UploadFile]) -> List[Dict[str, Any]]:
        session = self.get_or_create(session_id)
        if not files or len(files) > self.settings.max_upload_files:
            raise HTTPException(413, "Upload file count exceeds limit")
        session_dir = contained_path(self.upload_root, session_id)
        session_dir.mkdir(parents=True, exist_ok=True)

        saved_files: List[Dict[str, Any]] = []
        for upload in files:
            file_id = str(uuid.uuid4())
            safe_name = Path((upload.filename or f"{file_id}.bin").replace("\\", "/")).name
            if Path(safe_name).suffix.lower() not in {".pdf", ".txt", ".csv", ".tsv", ".json", ".docx", ".xlsx"}:
                raise HTTPException(415, "Unsupported document format")
            stored_name = f"{file_id}_{safe_name}"
            path = session_dir / stored_name
            size = 0
            try:
                with path.open("xb") as out_file:
                    while chunk := await upload.read(1024 * 1024):
                        size += len(chunk)
                        if size > self.settings.max_upload_mb * 1024 * 1024:
                            raise HTTPException(413, "Upload exceeds byte limit")
                        out_file.write(chunk)
            except Exception:
                path.unlink(missing_ok=True)
                raise

            info = {
                "file_id": file_id,
                "filename": stored_name,
                "display_name": safe_name,
                "path": str(path),
                "size": path.stat().st_size,
                "upload_time": datetime.now(timezone.utc).isoformat(),
                "status": "uploaded",
            }
            session["uploaded_files"].append(info)
            saved_files.append(info)
            logger.info(
                "file.saved session_id=%s display_name=%s size=%d path=%s",
                session_id,
                safe_name,
                info["size"],
                path,
            )

        self.save()
        return saved_files

    def get_files(self, session_id: str) -> List[Dict[str, Any]]:
        session = self.get_or_create(session_id)
        return session.get("uploaded_files", [])

    def find_file(self, session_id: str, filename: str) -> Optional[Path]:
        for file_info in self.get_files(session_id):
            candidates = [
                Path(file_info.get("path", "")),
                self.upload_root / session_id / file_info.get("filename", ""),
                self.upload_root / session_id / filename,
            ]
            for candidate in candidates:
                if candidate.resolve().is_relative_to(self.upload_root.resolve()) and candidate.is_file() and (candidate.name == filename or candidate.name.endswith(filename)):
                    return candidate
        return None

    def find_file_by_name(self, filename: str) -> Optional[Path]:
        for session in self.sessions.values():
            if session.get("user_id") != principal_id.get():
                continue
            for file_info in session.get("uploaded_files", []):
                path = Path(file_info.get("path", ""))
                if path.resolve().is_relative_to(self.upload_root.resolve()) and path.is_file() and (path.name == filename or path.name.endswith(filename)):
                    return path
        return None

    def table_state(self, session_id: str) -> Dict[str, Any]:
        session = self.get_or_create(session_id)
        return session.get("table_data", {"columns": [], "rows": []})

    def update_table_rows(self, session_id: str, rows: List[Dict[str, Any]]) -> None:
        session = self.get_or_create(session_id)
        columns = [{"id": "Document Name", "label": "Document Name", "dataType": "file"}]
        for row in rows:
            for key in row.keys():
                if key not in {"Document Name", "__evidence__"} and not any(col["label"] == key for col in columns):
                    columns.append({"id": key, "label": key, "dataType": "text"})
        session["table_data"] = {"columns": columns, "rows": rows}
        self.save()

    def extract_texts(
        self,
        session_id: str,
        query: str = "",
        fields: Optional[List[str]] = None,
        table_context: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, str]]:
        documents: List[Dict[str, str]] = []
        retrieval_query = self._retrieval_query(query, fields or [], table_context or self.table_state(session_id))
        for file_info in self.get_files(session_id):
            path = Path(file_info.get("path", ""))
            text = ""
            pages = []
            if path.suffix.lower() == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                if len(reader.pages) > self.settings.max_document_pages:
                    raise HTTPException(413, "Document page count exceeds limit")
                pages = [{"page_number": index, "text": page.extract_text() or ""} for index, page in enumerate(reader.pages, 1)]
                text = "\n".join(page["text"] for page in pages)
            elif path.suffix.lower() == ".docx":
                from docx import Document
                document = Document(str(path))
                text = "\n".join([paragraph.text for paragraph in document.paragraphs] + ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows])
            elif path.suffix.lower() == ".xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    text = "\n".join(f"Sheet: {sheet.title}\n" + "\n".join("\t".join(str(value) if value is not None else "" for value in row) for row in sheet.iter_rows(values_only=True)) for sheet in workbook)
                finally:
                    workbook.close()
            else:
                text = self._read_text(path)
            if not text.strip():
                raise HTTPException(422, "No readable text; scanned PDFs require OCR before extraction")
            if len(text) > self.settings.max_document_chars:
                raise HTTPException(413, "Document exceeds extraction context budget; split the document")
            text_index_path = self._write_text_index(session_id, file_info["file_id"], file_info["display_name"], text)
            embedding_path = self.embedding_index.ensure_index(
                session_id,
                file_info["file_id"],
                file_info["display_name"],
                text,
            )
            retrieved_chunks = self.embedding_index.retrieve(embedding_path, retrieval_query, top_k=14)
            retrieval_context_path = self.embedding_index.write_context_file(
                session_id,
                file_info["file_id"],
                file_info["display_name"],
                str(path),
                str(text_index_path),
                retrieval_query,
                retrieved_chunks,
            )
            max_chars = max(self.settings.max_document_chars, 0)
            selected_text = text if max_chars == 0 else text[:max_chars]
            documents.append({
                "file_id": file_info["file_id"],
                "file_name": file_info["display_name"],
                "path": str(path),
                "text_path": str(text_index_path),
                "embedding_index_path": str(embedding_path),
                "retrieval_context_path": str(retrieval_context_path),
                "text": selected_text,
                "pages": pages,
            })
            logger.info(
                "file.text_extracted session_id=%s file_name=%s chars=%d provided_chars=%d text_index=%s embedding_index=%s retrieved_chunks=%d",
                session_id,
                file_info["display_name"],
                len(text),
                len(selected_text),
                text_index_path,
                embedding_path,
                len(retrieved_chunks),
            )
        return documents

    def supporting_citation(self, session_id: str, file_name: str, field: str, value: Any) -> Dict[str, str]:
        excerpt = self.supporting_excerpt(session_id, file_name, field, value)
        return {
            "source": "file",
            "filename": file_name,
            "domain": file_name,
            "excerpt": excerpt,
            "field": str(field or ""),
            "value": str(value or ""),
        }

    def supporting_excerpt(self, session_id: str, file_name: str, field: str, value: Any) -> str:
        text = self._text_for_display_name(session_id, file_name)
        if not text:
            return "No extractable source text was found for this file."

        candidates = self._citation_search_terms(field, value)
        lower_text = text.lower()
        for term in candidates:
            if not term:
                continue
            index = lower_text.find(term.lower())
            if index >= 0:
                start = max(index - 500, 0)
                end = min(index + max(len(term), 1) + 900, len(text))
                return " ".join(text[start:end].split())

        retrieved = self._retrieved_supporting_chunk(session_id, file_name, f"{field} {value}")
        if retrieved:
            return retrieved
        return "No exact supporting chunk was found in the local text index for this value."

    def _retrieved_supporting_chunk(self, session_id: str, file_name: str, query: str) -> str:
        normalized_name = self._normalize_filename(file_name)
        for file_info in self.get_files(session_id):
            display_name = file_info.get("display_name") or file_info.get("filename") or ""
            if self._normalize_filename(display_name) != normalized_name:
                continue
            index_path = self.embedding_index.root / session_id / f"{file_info['file_id']}.json"
            chunks = self.embedding_index.retrieve(index_path, query, top_k=1)
            if chunks:
                return " ".join(str(chunks[0].get("text") or "").split())
        return ""

    def _text_for_display_name(self, session_id: str, file_name: str) -> str:
        normalized_name = self._normalize_filename(file_name)
        for file_info in self.get_files(session_id):
            display_name = file_info.get("display_name") or file_info.get("filename") or ""
            if self._normalize_filename(display_name) != normalized_name:
                continue
            text_index = self.index_root / session_id / f"{file_info['file_id']}_{Path(display_name).stem}.txt"
            if text_index.exists():
                return text_index.read_text(encoding="utf-8", errors="ignore")
            path = Path(file_info.get("path", ""))
            if path.suffix.lower() == ".pdf":
                return self._read_pdf(path)
            return self._read_text(path)
        return ""

    def _citation_search_terms(self, field: str, value: Any) -> List[str]:
        terms: List[str] = []
        raw_value = str(value or "").strip()
        raw_field = str(field or "").strip()
        if raw_value and raw_value.lower() != "not found":
            terms.append(raw_value)
            cleaned_value = re.sub(r"[$,]", "", raw_value)
            if cleaned_value != raw_value:
                terms.append(cleaned_value)
            number_parts = re.findall(r"\d+(?:\.\d+)?", raw_value)
            terms.extend(number_parts)
        if raw_field:
            terms.append(raw_field)
            words = [word for word in re.findall(r"[A-Za-z]+", raw_field) if len(word) > 3]
            if len(words) >= 2:
                terms.append(" ".join(words[:3]))
            terms.extend(words)
        seen = set()
        unique = []
        for term in terms:
            normalized = term.lower().strip()
            if normalized and normalized not in seen:
                seen.add(normalized)
                unique.append(term.strip())
        return unique

    def _normalize_filename(self, value: str) -> str:
        return re.sub(r"^[a-f0-9-]+_", "", str(value or "")).strip().lower()

    def _write_text_index(self, session_id: str, file_id: str, display_name: str, text: str) -> Path:
        session_dir = self.index_root / session_id
        session_dir.mkdir(parents=True, exist_ok=True)
        safe_name = Path(display_name).stem or file_id
        index_path = session_dir / f"{file_id}_{safe_name}.txt"
        index_path.write_text(text or "", encoding="utf-8", errors="ignore")
        return index_path

    def _retrieval_query(self, query: str, fields: List[str], table_context: Dict[str, Any]) -> str:
        columns = table_context.get("columns") if isinstance(table_context, dict) else []
        rows = table_context.get("rows") if isinstance(table_context, dict) else []
        column_text = " ".join(
            str((column.get("label") or column.get("id")) if isinstance(column, dict) else column)
            for column in (columns or [])
            if str(column.get("label") if isinstance(column, dict) else column) != "+"
        )
        row_text = ""
        if isinstance(rows, list):
            row_text = " ".join(
                " ".join(str(value) for value in row.values() if value not in (None, ""))
                for row in rows[:8]
                if isinstance(row, dict)
            )
        return " ".join(
            part for part in [
                str(query or ""),
                " ".join(str(field) for field in fields),
                column_text,
                row_text,
            ]
            if part
        )[:6000]

    def _read_text(self, path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return ""

    def _read_pdf(self, path: Path) -> str:
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            return ""
